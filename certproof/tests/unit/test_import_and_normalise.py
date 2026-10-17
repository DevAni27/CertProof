import io

import pytest
from openpyxl import Workbook

from backend.app.services.imports import tabular
from backend.app.services.verification.common import norm_cmp, squash


def test_normalisation_preserves_meaningful_characters():
    assert norm_cmp("Zoë  O'Brien") == "Zoë O'Brien"            # NFC, whitespace collapsed
    assert norm_cmp("A B\n C") == "A B C"                              # nbsp + line wrap
    assert norm_cmp("José") != norm_cmp("Jose")                             # accents are never stripped
    assert norm_cmp("anna") != norm_cmp("Anna")                             # case preserved
    assert norm_cmp("O'Brien") != norm_cmp("OBrien")                        # punctuation preserved
    assert squash("A B\nC") == "ABC"


def test_csv_preserves_leading_zero_ids_and_diacritics():
    data = "Name,ID\nJosé Núñez,007\n".encode("utf-8-sig")
    s = tabular.parse_csv(data, "x.csv", 1, 100)
    assert s.rows[0].cells == {"Name": "José Núñez", "ID": "007"}
    assert s.rows[0].row_index == 2


def test_csv_non_utf8_is_rejected_not_guessed():
    with pytest.raises(tabular.ImportProblem, match="UTF-8"):
        tabular.parse_csv("Name\nJosé".encode("latin-1"), "x.csv", 1, 100)


def test_csv_row_limit():
    data = ("Name\n" + "\n".join(f"n{i}" for i in range(11))).encode()
    with pytest.raises(tabular.ImportProblem, match="limit"):
        tabular.parse_csv(data, "x.csv", 1, 10)


def _xlsx(build):
    wb = Workbook()
    build(wb)
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def test_xlsx_sheet_selection_header_row_and_numeric_id_warning():
    def build(wb):
        ws = wb.active
        ws.title = "A"
        ws.append(["junk title row"])
        ws.append(["Name", "Student ID"])
        ws.append(["Ann", 7])          # numeric cell: leading zeros already lost
        wb.create_sheet("B").append(["Name"])
    d = _xlsx(build)
    s = tabular.parse_xlsx(d, "f.xlsx", "A", 2, 100)
    assert s.sheets == ["A", "B"] and s.headers == ["Name", "Student ID"]
    assert s.rows[0].row_index == 3
    assert any("leading zeros" in w for w in s.warnings)
    assert tabular.parse_xlsx(d, "f.xlsx", "B", 1, 100).rows == []


def test_xlsx_formula_without_cached_value_needs_review():
    def build(wb):
        ws = wb.active
        ws.append(["Name"])
        ws["A2"] = '="Ann"&" Lee"'       # openpyxl writes no cached value
    s = tabular.parse_xlsx(_xlsx(build), "f.xlsx", None, 1, 100)
    assert "FORMULA_NO_CACHED_VALUE" in s.rows[0].findings
    assert s.rows[0].cells["Name"] == ""     # formulas are never evaluated
    assert any("never executed" in w for w in s.warnings)


def test_macro_and_legacy_excel_rejected():
    for n in ("a.xlsm", "a.xls", "a.xlsb"):
        with pytest.raises(tabular.ImportProblem):
            tabular.parse_upload(b"x", n, None, 1, 10)


def test_duplicate_and_blank_headers_are_made_unique():
    s = tabular.parse_csv(b"Name,,Name\nA,B,C\n", "x.csv", 1, 10)
    assert s.headers == ["Name", "Column 2", "Name (2)"]
